"""2027 사업계획 시트(계획 뷰, view="plan27") + 양식 변경 자동 대응 — 2026-10-29.

합성 xlsx(openpyxl)로 실제 시트와 같은 모양(머리글 11~12행, 데이터 13행~, 매출행+원가행 2행, 맨 아래 합계 블록)을 만든다.
운영 data/ 파일은 안 건드림 — 상태·학습 맵·기준 스냅샷 경로를 전부 tmp_path로 바꿔치기한다.
"""
import json
import os

import pandas as pd
import pytest
from flask import Flask
from openpyxl import Workbook

import paths
import perf_autoupdate as pa
import performance as P

N_PAIRS = 60   # validate()가 매출행 50개 이상을 요구
BLOCK_LABELS = ["매출", "원가", "공통원가", "인건비", "자산", "관리비", "손익"]


def _make_plan_workbook(path, sheet="2027년 (사업계획)", shift_at=None, break_sum=False, with_block=True):
    """실제 2027 시트와 같은 배치의 합성 파일. shift_at: 그 열 번호 앞에 새 열을 끼워 이후를 한 칸 밀어냄(양식 변경 흉내).
    break_sum: 27 계획 합계를 1~12월 합과 어긋나게(계산식 검증 실패용)."""
    idx = {name: i for i, name in P._PERF_COL_MAP_PLAN27.items()}   # 필드 이름 → 시트 열 번호
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    col_of = lambda i: i + 1 if shift_at is not None and i >= shift_at else i   # noqa: E731  새 위치(0부터)

    def put(row, logical_idx, value):
        ws.cell(row=row, column=col_of(logical_idx) + 1, value=value)

    for i in range(0, 89):                       # 머리글 12행 — 열마다 고유한 글자(H<번호>)
        put(12, i, f"H{i}")
    if shift_at is not None:
        ws.cell(row=12, column=shift_at + 1, value="신규열")

    r = 13
    for p in range(N_PAIRS):
        part = "① 파트가" if p % 2 == 0 else "② 파트나"
        team = "SW팀" if p % 3 else "기획팀"
        common = {"team": team, "part": part, "use_yn": "사용", "progress": "착수",
                  "project_code": f"E{p:010d}01", "project_name": f"프로젝트{p}", "manager": "담당"}
        rev = dict(common, category="매출", plan_initial=100.0, jun_check_total=120.0 + (7 if break_sum and p == 0 else 0),
                   plan_diff_amount=20.0, profit_gross=70.0, cost_direct=50.0, cost_labor=10.0, cost_overhead=5.0, cost_mgmt=8.0,
                   operating_profit=47.0, profit_rate_raw=47 / 120)
        cost = dict(common, category="원가", plan_initial=40.0, jun_check_total=50.0, plan_diff_amount=10.0)
        for rec, months in ((rev, [10.0] * 12), (cost, [4.0] * 10 + [5.0, 5.0])):
            for name, val in rec.items():
                put(r, idx[name], val)
            for m, v in enumerate(months):
                put(r, idx[f"chk_m{m + 1:02d}"], v)
            r += 1
    if with_block:                               # 맨 아래 합계 블록 — 라벨은 파트 열, 값은 27 계획(jun_check_total)·26 추정(plan_initial) 열
        n = N_PAIRS
        y27 = [120 * n, 50 * n, 5 * n, 10 * n, 7, 8 * n, 47 * n]
        y26 = [100 * n, 40 * n, 6 * n, 9 * n, 0, 7 * n, 38 * n]
        r += 1
        for label, v27, v26 in zip(BLOCK_LABELS, y27, y26):
            put(r, idx["part"], label)
            put(r, idx["jun_check_total"], float(v27))
            put(r, idx["plan_initial"], float(v26))
            r += 1
    wb.save(path)


@pytest.fixture
def isolated_status(tmp_path, monkeypatch):
    """상태·학습 맵·기준 스냅샷을 tmp로 — 운영 data/perf_*.json을 건드리지 않게"""
    monkeypatch.setattr(pa, "_DATA_DIR", str(tmp_path))
    monkeypatch.setattr(pa, "STATUS_PATH", str(tmp_path / "status.json"))
    monkeypatch.setattr(pa, "REF_PATH", str(tmp_path / "ref.json"))
    monkeypatch.setattr(P, "PERF_LEARNED_MAPS_PATH", str(tmp_path / "learned.json"))
    # 열이 밀린 시트에서 H-chat을 부르지 않게(네트워크) — 못 찾은 열이 없으면 원래도 안 부름
    monkeypatch.setattr(pa, "ai_summary", lambda *a, **k: "")
    monkeypatch.setattr(pa, "ai_fill_missing", lambda *a, **k: {})
    return tmp_path


class TestPlanSheetName:
    def test_matches_plan_sheets_and_ignores_helper_sheets(self):
        assert P.is_plan_sheet("2027년 (사업계획)")
        assert P.is_plan_sheet("2028년 사업계획")
        assert not P.is_plan_sheet("(사업계획 참조)")                 # 연도 없는 보조 시트
        assert not P.is_plan_sheet("기타. 사업계획 재정리(260818)")
        assert not P.is_plan_sheet("2026년 (9월 집계)")               # 실적 시트
        assert not P.is_plan_sheet("참고1. 분류기준")

    def test_resolve_plan_sheet_picks_latest_year_or_none(self):
        sheets = ["2026년 (9월 집계)", "2027년 (사업계획)", "2028년 (사업계획)", "(사업계획 참조)"]
        assert P._resolve_plan_sheet(sheets) == "2028년 (사업계획)"
        assert P._resolve_plan_sheet(["2026년 (9월 집계)", "메모"]) is None

    def test_actual_resolver_never_picks_a_plan_sheet(self):
        assert P._resolve_perf_sheet(["2027년 (사업계획)", "2026년 (9월 집계)", "2026년 (8월 추정)"]) == "2026년 (9월 집계)"
        with pytest.raises(ValueError):
            P._resolve_perf_sheet(["2027년 (사업계획)"])

    def test_old_actual_maps_are_kept(self):
        """양식이 예전으로 되돌아와도 읽히게 — 지난 시트 맵을 지우지 않았는지"""
        for name in ("2026년 (6월 집계)", "2026년 (7월 추정)", "2026년 (8월 추정)", "2026년 (8월 집계)",
                     "2026년 (9월 추정)", "2026년 (9월 집계)", "2027년 (사업계획)"):
            assert P.perf_col_map(name), name


class TestFileOrdering:
    def test_date_at_end_of_name_is_recognised(self, tmp_path, monkeypatch):
        """…취합 완료_261028.xlsx — 날짜가 끝에 있어도 수정 시각 폴백이 아니라 날짜로 정렬(예전엔 옛 파일보다 뒤로 밀림)"""
        monkeypatch.setattr(paths, "DATA_DIR", str(tmp_path))
        old = tmp_path / "26년 사업계획 통합관리 파일_ver9.10_261020_9월 사업현황 리뷰.xlsx"
        new = tmp_path / "27년 사업계획 통합관리 파일_ver3.1_취합 완료_261028.xlsx"
        new.write_bytes(b"x")
        old.write_bytes(b"x")   # 수정 시각은 old가 더 최근
        assert os.path.basename(paths.perf_excel_candidates()[0]) == new.name


class TestReadPlanFrame:
    def test_reads_with_code_map_and_passes_validation(self, tmp_path):
        f = tmp_path / "plan.xlsx"
        _make_plan_workbook(f)
        df, sheet, year, month = P.read_perf_frame(str(f), view="plan27")
        assert (sheet, year, month) == ("2027년 (사업계획)", 2027, 0)
        rev = df[df["category"] == "매출"]
        assert len(rev) == N_PAIRS
        assert rev["jun_check_total"].sum() == 120 * N_PAIRS and rev["plan_initial"].sum() == 100 * N_PAIRS
        # 실적 전용 필드는 0 — 계획 시트엔 기준월·누계가 없음
        assert (df["jun_actual"] == 0).all() and (df["acc_operating_profit"] == 0).all()
        assert all(c["ok"] for c in pa.validate(df, "plan27"))

    def test_sheet_name_alone_selects_plan_mode(self, tmp_path):
        f = tmp_path / "plan.xlsx"
        _make_plan_workbook(f)
        _, _, _, month = P.read_perf_frame(str(f), "2027년 (사업계획)")   # view 생략 — 이름으로 판단
        assert month == 0

    def test_block_values_read_even_though_block_sums_include_unused_rows(self, tmp_path):
        """합계 블록은 필터 전 전체 행 합과 맞춤 — 사용여부가 빈 행이 있어도 블록을 믿을 수 있어야 함(2027 실파일: 매출 387.59 vs 사용 행 387.49억)"""
        f = tmp_path / "plan.xlsx"
        _make_plan_workbook(f)
        df, *_ = P.read_perf_frame(str(f), view="plan27")
        b27, b26 = df.attrs["plan_blocks"]["y27"], df.attrs["plan_blocks"]["y26"]
        assert b27 == {"overhead": 5.0 * N_PAIRS, "mgmt": 8.0 * N_PAIRS, "profit": 47.0 * N_PAIRS, "labor": 10.0 * N_PAIRS}
        assert b26 == {"overhead": 6.0 * N_PAIRS, "mgmt": 7.0 * N_PAIRS, "profit": 38.0 * N_PAIRS, "labor": 9.0 * N_PAIRS}

    def test_block_is_none_when_sheet_has_no_block(self, tmp_path):
        f = tmp_path / "plan.xlsx"
        _make_plan_workbook(f, with_block=False)
        df, *_ = P.read_perf_frame(str(f), view="plan27")
        assert df.attrs["plan_blocks"] == {"y27": None, "y26": None}


class TestPrepareFollowsFormatChanges:
    """사용자 요구: 양식이 또 바뀌어도 사람 손 없이 따라가고, 못 따라가면 틀린 숫자 대신 멈춘다"""

    def test_known_sheet_uses_existing_map_and_saves_reference_snapshot(self, isolated_status, tmp_path):
        f = tmp_path / "a.xlsx"
        _make_plan_workbook(f)
        df, sheet, year, month, st = pa.prepare(str(f), None, view="plan27")
        assert st["ok"] and st["method"] == "기존 매핑" and st["base"] == {"year": 2027}
        snap = json.loads((isolated_status / "ref.json").read_text(encoding="utf-8"))
        assert snap["plan27"]["sheet"] == "2027년 (사업계획)" and snap["plan27"]["labels"]

    def test_shifted_new_sheet_is_aligned_from_saved_snapshot(self, isolated_status, tmp_path):
        """새 시트(이름도 처음 보는 '2028년 (사업계획)', R열 앞에 열이 끼어 한 칸 밀림) — 직전 파일이 data/에 없어도(스냅샷만으로) 자동으로 맞춤"""
        first = tmp_path / "first.xlsx"
        _make_plan_workbook(first)
        pa.prepare(str(first), None, view="plan27")
        first.unlink()                                # 직전 파일을 치워도(경영실적 히스토리로 이동) 기준은 남아 있어야 함
        newf = tmp_path / "second.xlsx"
        _make_plan_workbook(newf, sheet="2028년 (사업계획)", shift_at=17)
        df, sheet, year, month, st = pa.prepare(str(newf), None, view="plan27")
        assert st["ok"] and st["method"].startswith("자동 정렬") and "저장된 머리글" in st["method"]
        assert all(c["ok"] for c in st["checks"])
        assert df[df["category"] == "매출"]["jun_check_total"].sum() == 120 * N_PAIRS
        learned = P.perf_col_map("2028년 (사업계획)")   # 학습한 맵이 저장돼 다음 읽기부턴 바로 쓰임
        assert learned[18] == "biz_type2" and learned[44 + 1] == "jun_check_total"

    def test_one_stray_row_is_tolerated_but_widespread_breakage_stops(self, isolated_status, tmp_path):
        """원본 수기 예외 한 줄(60행 중 1행 = 98.3%)은 95% 기준으로 봐주고, 널리 어긋나면 틀린 숫자를 내보내는 대신 멈춤"""
        one = tmp_path / "one.xlsx"
        _make_plan_workbook(one, break_sum=True)
        assert pa.prepare(str(one), None, view="plan27")[4]["ok"]

    def test_broken_numbers_stop_with_reason_and_keep_other_view_status(self, isolated_status, tmp_path):
        f = tmp_path / "bad.xlsx"
        _make_plan_workbook(f)
        from openpyxl import load_workbook
        wb = load_workbook(f)
        ws = wb.active
        sum_col = {v: k for k, v in P._PERF_COL_MAP_PLAN27.items()}["jun_check_total"] + 1
        for row in range(13, 13 + N_PAIRS * 2, 2):      # 매출행 전부의 27 계획 합계를 1~12월 합과 어긋나게
            ws.cell(row=row, column=sum_col, value=999.0)
        wb.save(f)
        pa.write_status({"ok": True, "sheet": "2026년 (9월 집계)", "file": "actual.xlsx"})   # 실적 상태가 이미 있음
        with pytest.raises(ValueError, match="계산식 검증 실패"):
            pa.prepare(str(f), None, view="plan27")
        plan = pa.read_status("plan27")
        assert plan["ok"] is False and "계산식 검증 실패" in plan["error"]
        actual = pa.read_status()
        assert actual["ok"] is True and actual["file"] == "actual.xlsx"      # 실적 상태는 그대로

    def test_file_without_plan_sheet_fails_clearly(self, isolated_status, tmp_path):
        f = tmp_path / "noplan.xlsx"
        wb = Workbook()
        wb.active.title = "2026년 (9월 집계)"
        wb.save(f)
        with pytest.raises(ValueError, match="사업계획 시트"):
            pa.prepare(str(f), None, view="plan27")


class TestStatusIsPerView:
    def test_actual_write_keeps_plan_status(self, isolated_status):
        pa.write_status({"ok": True, "serving": "p.xlsx"}, "plan27")
        pa.write_status({"ok": False, "error": "x"})                         # 실적 쪽이 통째로 덮어써도
        assert pa.read_status("plan27") == {"ok": True, "serving": "p.xlsx"}   # 계획 상태는 디스크 값 유지
        assert pa.read_status()["ok"] is False

    def test_plan_write_keeps_actual_status(self, isolated_status):
        pa.write_status({"ok": True, "sheet": "2026년 (9월 집계)"})
        pa.write_status({"ok": False, "error": "y"}, "plan27")
        assert pa.read_status()["sheet"] == "2026년 (9월 집계)"
        assert pa.read_status("plan27")["error"] == "y"

    def test_stale_plan_copy_inside_actual_write_is_ignored(self, isolated_status):
        """실적 코드가 읽어 둔 status(옛 plan27 포함)를 되쓰기 전에 계획 쪽이 갱신됐으면 새 값이 남아야 함"""
        pa.write_status({"ok": True, "serving": "old"}, "plan27")
        held = pa.read_status()                      # 실적 쪽이 읽어 둔 사본(plan27=old)
        pa.write_status({"ok": True, "serving": "new"}, "plan27")
        held["good"] = {"file": "f"}
        pa.write_status(held)
        assert pa.read_status("plan27")["serving"] == "new"


class TestPlanSummaryEndpoint:
    @pytest.fixture
    def client(self, tmp_path, monkeypatch):
        f = tmp_path / "plan.xlsx"
        _make_plan_workbook(f)
        df, sheet, year, _ = P.read_perf_frame(str(f), view="plan27")
        monkeypatch.setattr(P, "get_perf_df", lambda view="actual": df)
        monkeypatch.setitem(P._plan_cache, "year", year)
        monkeypatch.setitem(P._plan_cache, "sheet", sheet)
        app = Flask(__name__)
        app.register_blueprint(P.perf_bp)
        return app.test_client()

    def test_totals_compare_27_plan_with_26_estimate(self, client):
        s = client.get("/api/performance/summary?view=plan27").get_json()
        t = s["total"]
        assert s["view"] == "plan27" and s["year"] == 2027 and "base" not in s
        assert t["y27"]["revenue"] == 120 * N_PAIRS and t["y26"]["revenue"] == 100 * N_PAIRS
        assert t["y27"]["cost"] == 50 * N_PAIRS and t["y26"]["cost"] == 40 * N_PAIRS
        # 매출이익 = 매출 − (직접원가 + 공통원가)
        assert t["y27"]["gross"] == (120 - (50 + 5)) * N_PAIRS
        assert t["y26"]["gross"] == (100 - (40 + 6)) * N_PAIRS                 # 26 공통원가는 합계 블록에서
        assert t["y27"]["profit"] == 47 * N_PAIRS and t["y26"]["profit"] == 38 * N_PAIRS
        assert t["rate"]["revenue"] == 120.0 and t["rate"]["profit"] == round(47 / 38 * 100, 1)

    def test_no_actual_only_fields(self, client):
        t = client.get("/api/performance/summary?view=plan27").get_json()["total"]
        for k in ("jun_actual", "achieve_rate", "est_progress_rate", "plan_initial"):
            assert k not in t

    def test_team_filter_blanks_26_values_that_only_exist_for_whole_company(self, client):
        t = client.get("/api/performance/summary?view=plan27&team=SW팀").get_json()["total"]
        assert t["y26"]["overhead"] is None and t["y26"]["profit"] is None and t["y26"]["gross"] is None
        assert t["rate"]["profit"] is None and t["rate"]["grossProfit"] is None
        assert t["y27"]["revenue"] > 0 and t["y26"]["revenue"] > 0              # 프로젝트 행 합은 필터 따라 줄어듦

    def test_by_part_and_monthly(self, client):
        s = client.get("/api/performance/summary?view=plan27").get_json()
        assert set(s["by_part"]) == {"① 파트가", "② 파트나"}
        assert sum(p["rev27"] for p in s["by_part"].values()) == 120 * N_PAIRS
        assert all(p["rev_rate"] == 120.0 for p in s["by_part"].values())   # 120 ÷ 100 — 프론트는 계산 없이 표시만
        assert len(s["monthly"]) == 12 and sum(m["revenue"] for m in s["monthly"]) == 120 * N_PAIRS

    def test_per_head_uses_current_headcount_like_the_actual_view(self, client, monkeypatch):
        """인당 = 27 계획 금액 ÷ 관리자 입력 현재 인원(실적 보기와 같은 값·같은 함수). 인원 없는 파트는 None"""
        monkeypatch.setattr(P.headcount, "read_headcount", lambda: {"parts": {"파트가": 4}, "plan_parts": {}})
        by_part = client.get("/api/performance/summary?view=plan27").get_json()["by_part"]
        a, b = by_part["① 파트가"], by_part["② 파트나"]
        # ① 파트가 = 매출행 30개 → 27 매출 3600, 관리비 240, 경상손익 1410 (천원)
        assert a["headcount"] == 4 and a["headcount_partial"] is False
        assert a["rev_per_head"] == 3600 / 4 and a["mgmt_per_head"] == 240 / 4 and a["profit_per_head"] == 1410 / 4
        assert b["headcount"] is None and b["rev_per_head"] is None and b["profit_per_head"] is None   # 인원 미입력 → 제외

    def test_per_head_blank_for_parts_spanning_teams_when_team_filtered(self, client, monkeypatch):
        monkeypatch.setattr(P.headcount, "read_headcount", lambda: {"parts": {"파트가": 4, "파트나": 5}, "plan_parts": {}})
        by_part = client.get("/api/performance/summary?view=plan27&team=SW팀").get_json()["by_part"]
        for st in by_part.values():   # 합성 데이터는 두 파트 모두 두 팀에 걸쳐 있음 → 팀 몫 금액 ÷ 파트 전체 인원이 되므로 비움
            assert st["headcount_partial"] is True and st["rev_per_head"] is None
        whole = client.get("/api/performance/summary?view=plan27").get_json()["by_part"]
        assert all(st["headcount_partial"] is False for st in whole.values())

    def test_unsupported_endpoints_say_so(self, client):
        assert client.get("/api/performance/insights?view=plan27").status_code == 400
        assert client.get("/api/performance/summary/breakdown?view=plan27").get_json()["available"] is False

    def test_detail_table_works_without_finance_comparison(self, client):
        d = client.get("/api/performance/data?view=plan27&page_size=5").get_json()
        assert d["total"] == N_PAIRS
        row = d["data"][0]
        assert row["jun_check_total"] == 120.0 and row["plan_initial"] == 100.0
        assert row["est_cost"] == 50.0 and row["fin_mismatch"] == []

    def test_unknown_view_falls_back_to_actual(self, tmp_path, monkeypatch):
        called = []
        monkeypatch.setattr(P, "get_perf_df", lambda *a, **k: called.append(a) or pd.DataFrame())
        app = Flask(__name__)
        app.register_blueprint(P.perf_bp)
        app.test_client().get("/api/performance/options?view=nonsense")
        assert called == [()]                          # 실적 뷰는 인자 없이 호출(기존 테스트 호환)


class TestViewsEndpoint:
    """전환 칩이 쓰는 /api/performance/views — 칩 이름은 실제 시트명(2026-10-29)"""

    def _client(self, monkeypatch, plan_df, actual_df, serving_sheet):
        monkeypatch.setattr(P, "get_perf_df", lambda view="actual": plan_df if view == "plan27" else actual_df)
        monkeypatch.setattr(pa, "read_status", lambda view="actual": {"serving_sheet": serving_sheet})
        app = Flask(__name__)
        app.register_blueprint(P.perf_bp)
        return app.test_client()

    def test_both_views_available_with_sheet_names(self, tmp_path, monkeypatch):
        f = tmp_path / "plan.xlsx"
        _make_plan_workbook(f)
        df, sheet, year, _ = P.read_perf_frame(str(f), view="plan27")
        monkeypatch.setitem(P._plan_cache, "sheet", sheet)
        monkeypatch.setitem(P._plan_cache, "year", year)
        v = self._client(monkeypatch, df, df, "2026년 (9월 집계)").get("/api/performance/views").get_json()
        assert v["actual"]["available"] and v["actual"]["sheet"] == "2026년 (9월 집계)"
        assert v["plan27"] == {"available": True, "sheet": "2027년 (사업계획)", "year": 2027}

    def test_plan_hidden_when_file_has_no_plan_sheet(self, monkeypatch):
        monkeypatch.setitem(P._plan_cache, "sheet", None)
        monkeypatch.setitem(P._plan_cache, "year", 0)
        v = self._client(monkeypatch, pd.DataFrame(), pd.DataFrame({"x": [1]}), "2026년 (9월 집계)").get("/api/performance/views").get_json()
        assert v["plan27"]["available"] is False and v["actual"]["available"] is True
