"""프로젝트 상세(실적 탭) 엑셀 다운로드 — 화면 표와 같은 행·같은 기본 컬럼, 1뎁스만(2026-10-14).

배지는 숫자 옆 글자로 — 표시 형식(number_format)에 붙여서 값은 숫자 그대로(합계·정렬 가능), 칸 색·글자색도 화면과 같은 뜻.
  빨강 "상이" = 재무 이력(완료 보고)과 매출/직접원가 1,000원 이상 다름 / 파랑 "착수↑" = 직접원가 ≥ 착수 보고 직접원가
  칸 메모(마우스 올림)에 화면 툴팁과 같은 금액 상세, 맨 끝 "확인 사항" 열에도 글자로.
  (재무 이력 2뎁스를 행 그룹으로 넣었다가 억지스러워 뺌 — 2026-10-14)
금액은 원 단위 정수(실적 데이터는 천원 → ×1000).
"""
from __future__ import annotations

import io
import json
import re
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from shared import strip_part_prefix

# 현대 9색 파생 — 화면과 같은 의미색
_NAVY    = "002C5F"
_RED     = "C42A0E"
_RED_BG  = "FBE3DD"
_BLUE    = "006E8A"
_BLUE_BG = "DCEFF5"
_HEAD_BG = "E4DCD3"   # Hyundai Sand
_EMPH_BG = "FFF7DC"   # 경상손익·손익률 강조(화면 --emphasis-col-bg)
_MUTED   = "8A8580"

_THIN = Side(style="thin", color="D9D2CA")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_WON = "#,##0"
_PCT = '0.0"%"'

# (헤더, 필드, 종류, 폭) — 화면 프로젝트 상세 기본 컬럼 순서. 종류: txt / won(천원→원) / pct / num
_MAIN_COLS = [
    ("NO.",               "_group_no",        "num", 6),
    ("프로젝트코드",       "project_code",     "txt", 20),
    ("진행",              "progress",         "txt", 8),
    ("프로젝트명",         "project_name",     "txt", 44),
    ("팀",                "team",             "txt", 22),
    ("파트",              "part",             "txt", 14),
    ("담당자",             "manager",          "txt", 9),
    ("사업계획 매출",      "plan_initial",     "won", 15),
    ("사업계획 원가",      "plan_cost",        "won", 15),
    ("당월 추정 매출",     "jun_check_total",  "won", 19),   # 배지 글자 자리만큼 넓게
    ("당월 추정 직접원가", "est_cost",         "won", 24),
    ("매출 이익",          "est_gross",        "won", 15),
    ("인건비",             "cost_labor",       "won", 13),
    ("공통원가",           "cost_overhead",    "won", 13),
    ("관리비",             "cost_mgmt",        "won", 13),
    ("경상손익",           "operating_profit", "won", 15),
    ("손익률",             "profit_rate",      "pct", 9),
    ("확인 사항",          "_check",           "txt", 30),
]
_EMPH_FIELDS = {"operating_profit", "profit_rate"}


def _num(v):
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _won_text(v) -> str:
    return "-" if v is None else f"{round(v):,}원"


def _diff_text(a, b) -> str:
    if a is None or b is None:
        return ""
    d = round(a - b)
    return f"· 차이: {'+' if d > 0 else ''}{d:,}원"


def _won(row, field):
    n = _num(row.get(field))
    return None if n is None else n * 1000


def _badges(row) -> dict[str, list[tuple[str, str, str]]]:
    """필드 → [(tone, label, 설명)] — 화면 PerformanceTable/columns.finBadges와 같은 규칙"""
    out: dict[str, list[tuple[str, str, str]]] = {}
    mismatch = row.get("fin_mismatch") or []
    for field, name in (("jun_check_total", "매출"), ("est_cost", "원가")):
        if field in mismatch:
            perf, fin = _won(row, field), _won(row, f"fin_{field}")
            out.setdefault(field, []).append(("danger", "상이", "\n".join(filter(None, [
                f"⚠ 재무 이력(완료 보고)과 {name} 금액 상이",
                f"· 실적현황: {_won_text(perf)}", f"· 재무 이력: {_won_text(fin)}", _diff_text(perf, fin)]))))
    over = row.get("fin_cost_over_start")
    if over is not None and not (isinstance(over, float) and pd.isna(over)) and bool(over):
        perf, st = _won(row, "est_cost"), _won(row, "fin_start_est_cost")
        out.setdefault("est_cost", []).append(("info", "착수↑", "\n".join(filter(None, [
            "착수 보고 직접원가보다 크거나 같음",
            f"· 실적현황: {_won_text(perf)}", f"· 재무 이력(착수): {_won_text(st)}", _diff_text(perf, st)]))))
    return out


# 프론트 utils/projectCode.extractRealCode·countFinanceHistory와 같은 규칙 — 프로젝트코드 옆 숫자 배지(재무 이력 건수)
_REAL_CODE_RE = re.compile(r"^[A-Z]\d{6,}")


def finance_code_counts(fin_df: pd.DataFrame) -> dict[str, int]:
    """/api/finance/codes와 같은 코드 → 건수"""
    if fin_df is None or fin_df.empty or "project_code" not in fin_df.columns:
        return {}
    s = fin_df["project_code"].astype(str).str.strip()
    return s[s.ne("") & s.ne("0")].value_counts().to_dict()


def _finance_count(code, counts: dict[str, int], cache: dict[str, int]) -> int:
    """앞쪽 정식 코드만 잘라 정확히 일치 → 없으면 부분일치 합. placeholder 코드는 0"""
    m = _REAL_CODE_RE.match(str(code or "").strip())
    if not m:
        return 0
    term = m.group(0)
    if term not in cache:
        n = counts.get(term, 0)
        if not n:
            low = term.lower()
            n = sum(c for k, c in counts.items() if low in str(k).lower())
        cache[term] = n
    return cache[term]


def _write_value(cell, kind: str, v):
    if kind in ("won", "pct", "num"):
        n = _num(v)
        cell.value = None if n is None else (round(n * 1000) if kind == "won" else n)
        cell.number_format = _WON if kind == "won" else _PCT if kind == "pct" else "0"
        cell.alignment = Alignment(horizontal="right", vertical="center")
    else:
        cell.value = "" if v is None or (isinstance(v, float) and pd.isna(v)) else str(v)
        cell.alignment = Alignment(vertical="center")


def build_project_detail_xlsx(rows: pd.DataFrame, condition: str, fin_counts: dict[str, int] | None = None) -> bytes:
    """fin_counts: 재무 이력 코드 → 건수(finance_code_counts) — 있으면 프로젝트코드 뒤에 " [N]"(화면 숫자 배지)"""
    wb = Workbook()
    ws = wb.active
    ws.title = "프로젝트 상세"

    # ── 머리말 3줄: 제목 · 조건 · 배지 범례 ──
    ws["A1"] = "프로젝트 상세"
    ws["A1"].font = Font(size=14, bold=True, color=_NAVY)
    ws["A2"] = condition
    ws["A2"].font = Font(size=9, color=_MUTED)
    ws["A3"] = "상이"
    ws["A3"].font = Font(bold=True, color=_RED)
    ws["A3"].fill = PatternFill("solid", fgColor=_RED_BG)
    ws["B3"] = "재무 이력(완료 보고)과 매출·직접원가가 1,000원 이상 다름"
    ws["E3"] = "착수↑"
    ws["E3"].font = Font(bold=True, color=_BLUE)
    ws["E3"].fill = PatternFill("solid", fgColor=_BLUE_BG)
    ws["F3"] = "당월 추정 직접원가 ≥ 재무 이력 착수 보고 직접원가"
    ws["J3"] = "배지 칸에 마우스를 올리면 금액 상세 · 프로젝트코드 뒤 [N] = 재무 보고서 건수"
    for c in ("B3", "F3", "J3"):
        ws[c].font = Font(size=9, color=_MUTED)

    head_row = 5
    for ci, (title, _, _, width) in enumerate(_MAIN_COLS, 1):
        c = ws.cell(head_row, ci, title)
        c.font = Font(bold=True, color=_NAVY)
        c.fill = PatternFill("solid", fgColor=_HEAD_BG)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = _BORDER
        ws.column_dimensions[get_column_letter(ci)].width = width
    ws.row_dimensions[head_row].height = 22
    ws.freeze_panes = ws.cell(head_row + 1, 3)

    col_of = {field: ci for ci, (_, field, _, _) in enumerate(_MAIN_COLS, 1)}
    count_cache: dict[str, int] = {}
    r = head_row
    for _, row in rows.iterrows():
        r += 1
        rec = row.to_dict()
        rec["part"] = strip_part_prefix(str(rec.get("part") or ""))
        n_fin = _finance_count(rec.get("project_code"), fin_counts or {}, count_cache)
        if n_fin:
            rec["project_code"] = f"{rec.get('project_code')} [{n_fin}]"
        badges = _badges(rec)
        rec["_check"] = " · ".join(
            f"{label}({'매출' if f == 'jun_check_total' else '직접원가'})"
            for f, items in badges.items() for _, label, _ in items)

        for ci, (_, field, kind, _) in enumerate(_MAIN_COLS, 1):
            c = ws.cell(r, ci)
            _write_value(c, kind, rec.get(field))
            c.border = _BORDER
            if field in _EMPH_FIELDS:
                c.fill = PatternFill("solid", fgColor=_EMPH_BG)

        # 배지 — 숫자 뒤에 배지 글자(표시 형식), 칸 색·글자색, 메모
        for field, items in badges.items():
            c = ws.cell(r, col_of[field])
            danger = any(t == "danger" for t, _, _ in items)
            labels = " ".join(f"[{label}]" for _, label, _ in items)
            c.number_format = f'{_WON}" {labels}"'
            c.fill = PatternFill("solid", fgColor=_RED_BG if danger else _BLUE_BG)
            c.font = Font(bold=True, color=_RED if danger else _BLUE)
            note = "\n\n".join(t for _, _, t in items)
            c.comment = Comment(note, "대시보드", width=280, height=40 + 16 * note.count("\n"))
        if rec["_check"]:
            danger = any(t == "danger" for items in badges.values() for t, _, _ in items)
            ws.cell(r, col_of["_check"]).font = Font(bold=True, color=_RED if danger else _BLUE)

    # 머리글 자동 필터(엑셀에서 열별로 거를 수 있게)
    if r > head_row:
        ws.auto_filter.ref = f"A{head_row}:{get_column_letter(len(_MAIN_COLS))}{r}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def condition_text(args, base: dict | None) -> str:
    parts = args.getlist("part")
    bits = [f"다운로드 {datetime.now():%Y-%m-%d %H:%M}"]
    if base:
        bits.append(f"기준 {base.get('year')}년 {base.get('month')}월 시트")
    bits.append(f"팀 {args.get('team') or '전체'}")
    bits.append(f"파트 {', '.join(strip_part_prefix(p) for p in parts) if parts else '전체'}")
    if args.get("progress"):
        bits.append(f"진행 {args.get('progress')}")
    if args.get("search"):
        bits.append(f"검색 \"{args.get('search')}\"")
    bits += _describe_cf(args.get("cf", ""))
    return " · ".join(bits)


def _describe_cf(raw: str) -> list[str]:
    """열 필터 → "매출 이익 ≥ 2.1억" 같은 글자(화면 칩과 같은 모양)"""
    try:
        specs = json.loads(raw) if raw else []
    except ValueError:
        return []
    titles = {field: title for title, field, _, _ in _MAIN_COLS}
    out = []
    for f in specs if isinstance(specs, list) else []:
        if not isinstance(f, dict):
            continue
        t = titles.get(f.get("id"), f.get("id"))
        if f.get("type") == "text" and f.get("q"):
            out.append(f"{t}: \"{f['q']}\"")
        elif f.get("type") == "select" and f.get("values"):
            out.append(f"{t}: {', '.join(map(str, f['values']))}")
        elif f.get("type") == "number":
            u = "억" if f.get("unit") == "eok" else "%"
            lo, hi = f.get("min"), f.get("max")
            if lo is not None and hi is not None:
                out.append(f"{t}: {lo}~{hi}{u}")
            elif lo is not None:
                out.append(f"{t} ≥ {lo}{u}")
            elif hi is not None:
                out.append(f"{t} ≤ {hi}{u}")
    return out
